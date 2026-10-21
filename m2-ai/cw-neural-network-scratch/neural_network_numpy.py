# excel demo
# https://docs.google.com/spreadsheets/d/1GABf2jSFXd5rvPxJSTVoMK_ooPmykxpk1KZDDZ9_Dbg/edit?usp=sharing
# validation sheet
# https://docs.google.com/spreadsheets/d/1rHpZ68F_UkG3CMKPJ43tV8r6wItKA_nJw_6Na29LRv8/edit?usp=sharing

import os
from scipy import stats
import numpy as np
import math
import matplotlib.pyplot as plt
import random
import pandas as pd
import openpyxl

sigmoid = (
  lambda x:1 / (1 + np.exp(-x)),
  lambda x:x * (1 - x)
  )

def derivada_relu(x):
  x[x<=0] = 0
  x[x>0] = 1
  return x

relu = (
  lambda x: x * (x > 0),
  lambda x:derivada_relu(x)
  )

def circulo(num_datos = 100,R = 1, minimo = 0,maximo= 1):
  pi = math.pi
  r = R * np.sqrt(stats.truncnorm.rvs(minimo, maximo, size= num_datos)) * 10
  theta = stats.truncnorm.rvs(minimo, maximo, size= num_datos) * 2 * pi *10
  #print(len(r))
  #print(len(theta))
  x = np.cos(theta) * r
  y = np.sin(theta) * r

  y = y.reshape((num_datos,1))
  x = x.reshape((num_datos,1))

  #Vamos a reducir el numero de elementos para que no cause un Overflow
  x = np.round(x,3)
  y = np.round(y,3)

  df = np.column_stack([x,y])
  #print(df)
  return(df)

#print(len(X))
#print(len(Y))
# Numero de neuronas en cada capa.
# El primer valor es el numero de columnas de la capa de entrada.
#neuronas = [2,4,8,1]

# Funciones de activacion usadas en cada capa.
#funciones_activacion = [relu,relu, sigmoid]

#red_neuronal = []

#for paso in range(len(neuronas)-1):
#  x = capa(neuronas[paso],neuronas[paso+1],funciones_activacion[paso])
#  print(neuronas[paso],neuronas[paso+1],funciones_activacion[paso])

#  red_neuronal.append(x)
#  print(x)

#print(red_neuronal)


def mse(Ypredich, Yreal):

  # Calculamos el error
  #print(Ypredich)
  #print(Yreal)
  x = (np.array(Ypredich) - np.array(Yreal)) ** 2
  #print(x)
  x = np.mean(x)
  #print(x)
  # Calculamos la derivada de la funcion
  y = np.array(Ypredich) - np.array(Yreal)
  return (x,y)


def entrenamiento(epoch, X,Y, red_neuronal, lr = 0.01):

  # Output guardara el resultado de cada capa
  # En la capa 1, el resultado es el valor de entrada
  output = [X]
  for num_capa in range(len(red_neuronal)):
    z = output[-1] @ red_neuronal[num_capa].W + red_neuronal[num_capa].b
    a = red_neuronal[num_capa].funcion_act[0](z)
    #print(epoch)
    #if (epoch==0) or (epoch==-999):
      #print(epoch)
      #print(red_neuronal[num_capa].funcion_act[0]( np.array([-1, 1]) ))
      #print(X, X.shape)
      #print('w', red_neuronal[num_capa].W.shape)
      #print(red_neuronal[num_capa].W)
      #print('b', red_neuronal[num_capa].b.shape)
      #print(red_neuronal[num_capa].b)

      #print('z', z.shape)
      #print(z)
      #print('act ', num_capa)
      #print(a)
    # Incluimos el resultado de la capa a output
    output.append(a)

  # Backpropagation

  back = list(range(len(output)-1))
  back.reverse()

  # Guardaremos el error de la capa en delta
  delta = []

  for capa in back:
    # Backprop #delta

    a = output[capa+1]
    #if (epoch==0) or (epoch==-999):
    #  print('activation', a)

    if capa == back[0]:
      x = mse(a,Y)[1] * red_neuronal[capa].funcion_act[1](a)
      delta.append(x)
      if (epoch==0) and (epoch==-999):
        print('mse 0', x)
        #print('mse 01', mse(a,Y)[1] )
        #print('mse 02', red_neuronal[capa].funcion_act[1](a) )
        #print('mse 02', red_neuronal[capa].funcion_act[1](np.array([-2, -1, 0, 1, 2])) )

    else:
      if (epoch==0) and (capa == back[1]) and (epoch==-999):
        print('function a ', a)
        print('delta', delta[-1])
        print('W_temp', W_temp)

      x = delta[-1] @ W_temp * red_neuronal[capa].funcion_act[1](a)
      if (epoch==0) and (capa == back[1]) and (epoch==-999):
        print('x ', x)

      delta.append(x)

 #       print('mse 2', x)

    W_temp = red_neuronal[capa].W.transpose()
    #if (epoch==0) or (epoch==-999):
    #  print('W_temp', capa, red_neuronal[capa].W, W_temp)
    # Gradient Descent #
    red_neuronal[capa].b = red_neuronal[capa].b - np.mean(delta[-1], axis = 0, keepdims = True) * lr
    red_neuronal[capa].W = red_neuronal[capa].W - output[capa].transpose() @ delta[-1] * lr
    #if (epoch==0) and (capa==0):
    #  print('output 0', output[capa].transpose())
    #  print('delta  0', delta[-1])

    #print(red_neuronal[capa].W)
    #print('b', capa)
    #print(red_neuronal[capa].b)
 
  return output[-1]

class capa():
  def __init__(self, n_neuronas_capa_anterior, n_neuronas, funcion_act):
    self.funcion_act = funcion_act
    self.b  = np.round(stats.truncnorm.rvs(-1, 1, loc=0, scale=1, size= n_neuronas).reshape(1,n_neuronas),3)
    self.W  = np.round(stats.truncnorm.rvs(-1, 1, loc=0, scale=1, size= n_neuronas * n_neuronas_capa_anterior).reshape(n_neuronas_capa_anterior,n_neuronas),3)
    #self.b  = np.full((1,n_neuronas), 0.25).reshape(1,n_neuronas)
    #self.W  = np.full((1,n_neuronas * n_neuronas_capa_anterior), 0.75).reshape(n_neuronas_capa_anterior,n_neuronas)


base_dir = os.path.dirname(os.path.abspath(__file__))
c1_path = os.path.join(base_dir, "data", "city1.csv")
c2_path = os.path.join(base_dir, "data", "city2.csv")

if not os.path.exists(c1_path) or not os.path.exists(c2_path):
  from dataset import save_datasets
  save_datasets(output_dir=os.path.join(base_dir, "data"), num_points=150)

N = 150

# datos_1 = circulo(num_datos = N, R = 2)
# datos_2 = circulo(num_datos = N, R = 0.5)
datos_1 = pd.read_csv(c1_path)[["lat", "lon"]].values[:N]
datos_2 = pd.read_csv(c2_path)[["lat", "lon"]].values[:N]
X = np.concatenate([datos_1,datos_2])
X = np.round(X,3)
#X = np.array([[ -2.848,   5.975],
# [-10.429,  -8.024],
# [ -1.284,   2.872],
# [  4.758,   5.955],
# [ -3.938,  -5.78 ],
# [  8.525,  17.866],
# [-13.715,   4.789],
# [ -3.19,   13.309],
# [  3.227,  18.344],
# [ -5.014, -17.507],
# [  0.491,   1.897],
# [ -0.54,    1.12 ],
# [ -2.032,   1.092],
# [ -1.476,  -1.606],
# [ -0.084,   3.386],
# [ -1.618,   1.16 ],
# [  2.238,  -1.671],
# [ -2.102,   1.924],
# [ -1.059,   4.297],
# [  2.642,   2.981]])

Y = [0] * N + [1] * N
Y = np.array(Y).reshape(len(Y),1)

np.random.seed(3)
neuronas = [2,4,8,1]
funciones_activacion = [relu,relu, sigmoid]
red_neuronal = []

for paso in list(range(len(neuronas)-1)):
  x = capa(neuronas[paso],neuronas[paso+1],funciones_activacion[paso])
  red_neuronal.append(x)
  #print('w')
  #print(neuronas[paso])
  #print(neuronas[paso+1])
  
  #print(x.W)
  #print('b')
  #print(x.b)
  #print('---')
error = []
predicciones = []

for epoch in range(0, 500):
  ronda = entrenamiento(epoch, X = X ,Y = Y ,red_neuronal = red_neuronal, lr = 0.001)
  #print('-----')
  #print(ronda)
  #print(Y)
  predicciones.append(ronda)
  temp = mse(np.round(predicciones[-1]),Y)[0]
  #print(temp)
  #print('-----')
  error.append(temp)

print('===final Y1 (first 10) =====')
print(predicciones[-1][0:10])
print('===final Y2 (first 10) =====')
print(predicciones[-1][N:N+10])
print('===final MSE =====')
print(error[-1])

print('=== Layer 0 (2 -> 4) W and b ===')
print('W0:\n', red_neuronal[0].W)
print('b0:\n', red_neuronal[0].b)

print('=== Layer 1 (4 -> 8) W and b ===')
print('W1:\n', red_neuronal[1].W)
print('b1:\n', red_neuronal[1].b)

print('=== Layer 2 (8 -> 1) W and b ===')
print('W2:\n', red_neuronal[2].W)
print('b2:\n', red_neuronal[2].b)

# Update Excel validation sheet (only loading 20 coordinates and W, b per layer)
template_path = os.path.join(base_dir, "template_red_neuronal.xlsx")
output_path = os.path.join(base_dir, "validation_weights.xlsx")
if os.path.exists(template_path):
  wb = openpyxl.load_workbook(template_path)
  ws = wb["tensorflow validation N=20"]

  X_val = np.vstack([datos_1[:10], datos_2[:10]])
  for i in range(20):
    row = 3 + i
    ws.cell(row=row, column=2, value=float(X_val[i, 0]))
    ws.cell(row=row, column=3, value=float(X_val[i, 1]))

  # Layer 0 (2 -> 4): W in E3:H4, b in E8:H8
  for r in range(2):
    for c in range(4):
      ws.cell(row=3 + r, column=5 + c, value=float(red_neuronal[0].W[r, c]))
  for c in range(4):
    ws.cell(row=8, column=5 + c, value=float(red_neuronal[0].b[0, c]))

  # Layer 1 (4 -> 8): W in T3:AA6, b in T9:AA9
  for r in range(4):
    for c in range(8):
      ws.cell(row=3 + r, column=20 + c, value=float(red_neuronal[1].W[r, c]))
  for c in range(8):
    ws.cell(row=9, column=20 + c, value=float(red_neuronal[1].b[0, c]))

  # Layer 2 (8 -> 1): W in AU3:AU10, b in AU13
  for r in range(8):
    ws.cell(row=3 + r, column=47, value=float(red_neuronal[2].W[r, 0]))
  ws.cell(row=13, column=47, value=float(red_neuronal[2].b[0, 0]))

  wb.save(output_path)
  print(f"Saved validation sheet to {output_path}")

#epoch = list(range(0,1))
#plt.cla()
#plt.plot(epoch, error)
